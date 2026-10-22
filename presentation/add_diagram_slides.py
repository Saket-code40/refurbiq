"""Insert the four draw.io-style diagrams into RefurbIQ_PBL3.pptx.

Adds four new slides (same 'Title Only' layout as the rest of the deck,
title run formatting cloned from an existing slide) and reorders them
into the flow:

  after 10 Proposed Methodology   -> Methodology: Pipeline Overview (diagram 3)
  after 13 Methodology: Models    -> Technology Stack               (diagram 1)
  after 15 Result: Aspect/Platform-> System Design: App Workflow     (diagram 2)
  after that                      -> Recommendation Engine           (diagram 4)

Run:  python presentation/add_diagram_slides.py
"""

from __future__ import annotations

import copy
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Inches

HERE = Path(__file__).resolve().parent
DECK = HERE / "RefurbIQ_PBL3.pptx"
DIAGRAMS = HERE.parent / "diagrams"

SLIDE_W = Emu(9144000).inches  # 10"
SLIDE_H = Emu(6858000).inches  # 7.5"
TOP = 1.24          # below the title bar
MAX_W = 9.15
MAX_H = 5.62        # stop above the slide-number strip (~6.9")

NEW_SLIDES = [
    {
        "title": "Methodology: Pipeline Overview",
        "image": DIAGRAMS / "3_diagram_research_pipeline.png",
        "notes": (
            "This is the research pipeline end to end. Step 1 was an honest "
            "verification pass: none of the five platforms exposes a structured "
            "public review corpus, so the modelling corpus is a clearly labelled "
            "synthetic demonstration dataset. The pipeline is separately validated "
            "on 2,000 real Amazon reviews, agreeing with human star ratings in "
            "75.6% of annotations."
        ),
    },
    {
        "title": "Technology Stack",
        "image": DIAGRAMS / "1_diagram_tech_stack.png",
        "notes": (
            "Everything is open source. Streamlit serves the eight-tab interface; "
            "Hugging Face transformers provide zero-shot RoBERTa and fine-tuned "
            "DistilBERT sentiment; scikit-learn and XGBoost give the TF-IDF "
            "baselines; pandas and matplotlib handle data and visualisation. "
            "Model weights stay local in models/ and are git-ignored."
        ),
    },
    {
        "title": "System Design: App Workflow",
        "image": DIAGRAMS / "2_diagram_app_workflow.png",
        "notes": (
            "The app takes an uploaded review CSV (or the bundled demo corpus), "
            "cleans it, extracts aspects through the 13-aspect keyword vocabulary, "
            "runs whole-review sentiment, and produces the ABSA results table. That "
            "feeds two branches: the analytics dashboard and the recommendation "
            "engine, which powers Compare Platforms, Best Place to Buy and Explain "
            "Recommendation."
        ),
    },
    {
        "title": "Recommendation Engine",
        "image": DIAGRAMS / "4_diagram_recommendation.png",
        "notes": (
            "The overall score is a weighted sum of six transparent sub-scores: "
            "sentiment 0.25, product quality 0.25, service 0.15, value 0.15, "
            "trust 0.10 and return/refund 0.10. A confidence gate requires at "
            "least 20 reviews before any recommendation is made, and HIGH "
            "confidence additionally needs 50+ reviews with consistent sentiment."
        ),
    },
]


def clone_title_format(src_title, dst_title):
    """Copy paragraph + run properties from an existing deck title."""
    src_p = src_title.text_frame.paragraphs[0]
    dst_p = dst_title.text_frame.paragraphs[0]
    if src_p._p.pPr is not None:
        dst_p._p.get_or_add_pPr()
        dst_p._p.replace(dst_p._p.pPr, copy.deepcopy(src_p._p.pPr))
    if src_p.runs and dst_p.runs and src_p.runs[0]._r.rPr is not None:
        dst_p.runs[0]._r.insert(0, copy.deepcopy(src_p.runs[0]._r.rPr))


def main():
    prs = Presentation(str(DECK))
    layout = next(l for l in prs.slide_layouts if l.name == "Title Only")
    src_title = prs.slides[9].shapes.title  # 'Proposed Methodology' title

    n_orig = len(prs.slides.__iter__.__self__._sldIdLst)
    new_slide_objs = []
    for spec in NEW_SLIDES:
        slide = prs.slides.add_slide(layout)
        title_ph = slide.shapes.title
        title_ph.text = spec["title"]
        clone_title_format(src_title, title_ph)

        w_px, h_px = Image.open(spec["image"]).size
        ar = w_px / h_px
        w = min(MAX_W, MAX_H * ar)
        h = w / ar
        left = Inches((SLIDE_W - w) / 2)
        top = Inches(TOP + (MAX_H - h) / 2)
        slide.shapes.add_picture(str(spec["image"]), left, top, width=Inches(w))

        slide.notes_slide.notes_text_frame.text = spec["notes"]
        new_slide_objs.append(slide)
        print(f"added: {spec['title']!r} ({w:.2f}x{h:.2f} in)")

    # Reorder: originals are indices 0..n_orig-1, new slides n_orig..n_orig+3
    # (pipeline, techstack, appworkflow, recommendation).
    p, t, a, r = n_orig, n_orig + 1, n_orig + 2, n_orig + 3
    order = [0, 1, 2, 3, 4, 5, 6, 7, 8, 9,
             p,              # Pipeline Overview after Proposed Methodology
             10, 11, 12,
             t,              # Technology Stack after Methodology: Models
             13, 14,
             a, r,           # App Workflow + Recommendation before Outcome
             15, 16]
    assert sorted(order) == list(range(n_orig + 4))
    sldIdLst = prs.slides._sldIdLst
    els = list(sldIdLst)
    for el in els:
        sldIdLst.remove(el)
    for idx in order:
        sldIdLst.append(els[idx])

    prs.save(str(DECK))
    print(f"saved {DECK} with {len(order)} slides")

    # verify
    check = Presentation(str(DECK))
    for i, s in enumerate(check.slides, 1):
        t = ""
        pics = 0
        for sh in s.shapes:
            if sh.shape_type == 13:
                pics += 1
            elif sh.has_text_frame and not t and sh.text_frame.text.strip():
                t = sh.text_frame.text.strip().splitlines()[0][:48]
        print(f"  {i:2d}: pics={pics} | {t}")


if __name__ == "__main__":
    main()